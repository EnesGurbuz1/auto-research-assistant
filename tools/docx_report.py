"""
DOCX Report Generator - Word rapor üretim aracı.
Hocanın şablonuna (rapor_template.docx) uygun formatta rapor üretir.
python-docx ve docxtpl kullanır.
"""

import re
from pathlib import Path
from typing import Dict, Optional
from loguru import logger

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    Document = None

try:
    from docxtpl import DocxTemplate
except ImportError:
    DocxTemplate = None


class DocxReport:
    """DOCX rapor üretim aracı."""

    def __init__(self, template_path: Optional[Path] = None):
        self.template_path = template_path
        
        if Document is None:
            logger.warning("python-docx yüklü değil: pip install python-docx")

    def generate_from_markdown(self, markdown_content: str, output_path: Path) -> Path:
        """Markdown içeriğini DOCX'e dönüştür."""
        if Document is None:
            logger.error("python-docx yüklü değil!")
            # Fallback: txt olarak kaydet
            output_path.with_suffix(".txt").write_text(markdown_content, encoding="utf-8")
            return output_path.with_suffix(".txt")

        doc = Document()
        
        # Şablon varsa kullan
        if self.template_path and self.template_path.exists():
            try:
                doc = Document(str(self.template_path))
                logger.info(f"Şablon yüklendi: {self.template_path}")
            except Exception as e:
                logger.warning(f"Şablon yüklenemedi, boş doküman kullanılıyor: {e}")
                doc = Document()

        # Markdown'ı parse et ve DOCX'e dönüştür
        self._parse_markdown_to_docx(doc, markdown_content)
        
        # Kaydet
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"DOCX kaydedildi: {output_path}")
        
        return output_path

    def generate_from_template(self, context: Dict, output_path: Path) -> Path:
        """Şablondaki placeholder'ları doldurarak rapor üret."""
        if DocxTemplate is None:
            logger.error("docxtpl yüklü değil: pip install docxtpl")
            return output_path

        if not self.template_path or not self.template_path.exists():
            logger.error(f"Şablon bulunamadı: {self.template_path}")
            return output_path

        try:
            tpl = DocxTemplate(str(self.template_path))
            tpl.render(context)
            
            output_path.parent.mkdir(parents=True, exist_ok=True)
            tpl.save(str(output_path))
            logger.info(f"Template-based DOCX üretildi: {output_path}")
        except Exception as e:
            logger.error(f"Template render hatası: {e}")

        return output_path

    def _has_style(self, doc, style_name: str) -> bool:
        """Doküman bu style'a sahip mi kontrol et."""
        try:
            doc.styles[style_name]
            return True
        except KeyError:
            return False

    def _add_styled_paragraph(self, doc, text: str, style_name: str):
        """Style varsa kullan, yoksa normal paragraf + prefix."""
        if self._has_style(doc, style_name):
            doc.add_paragraph(text, style=style_name)
        elif style_name == "List Bullet":
            doc.add_paragraph(f"• {text}", style="List Paragraph" if self._has_style(doc, "List Paragraph") else "Normal")
        elif style_name == "List Number":
            doc.add_paragraph(text, style="List Paragraph" if self._has_style(doc, "List Paragraph") else "Normal")
        else:
            doc.add_paragraph(text)

    def _parse_markdown_to_docx(self, doc: Document, content: str):
        """Basit Markdown parser → DOCX dönüştürücü."""
        lines = content.split("\n")
        
        for line in lines:
            stripped = line.strip()
            
            if not stripped:
                continue
            
            # Heading'ler
            if stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("#### "):
                doc.add_heading(stripped[5:], level=4)
            
            # Bullet list
            elif stripped.startswith("- ") or stripped.startswith("* "):
                text = stripped[2:]
                text = self._clean_markdown_formatting(text)
                self._add_styled_paragraph(doc, text, "List Bullet")
            
            # Numbered list
            elif re.match(r"^\d+\.\s", stripped):
                text = re.sub(r"^\d+\.\s", "", stripped)
                text = self._clean_markdown_formatting(text)
                self._add_styled_paragraph(doc, text, "List Number")
            
            # Horizontal rule
            elif stripped in ("---", "***", "___"):
                doc.add_paragraph("─" * 50)
            
            # Normal paragraf
            else:
                text = self._clean_markdown_formatting(stripped)
                doc.add_paragraph(text)

    def _clean_markdown_formatting(self, text: str) -> str:
        """Markdown formatlamasını temizle (bold, italic vb.)."""
        # Bold
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        # Italic
        text = re.sub(r"\*(.+?)\*", r"\1", text)
        # Inline code
        text = re.sub(r"`(.+?)`", r"\1", text)
        # Links
        text = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", text)
        
        return text
