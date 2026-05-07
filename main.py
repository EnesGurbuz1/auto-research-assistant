#!/usr/bin/env python3
"""
Tez Araştırma Asistanı - Ana Giriş Noktası
Multi-Agent Grid Load Negotiation for EV Charging
"""

import sys
import os
import json
import click
import yaml
from pathlib import Path
from datetime import datetime
from dotenv import load_dotenv
from loguru import logger
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

# Proje kök dizinini ayarla
PROJECT_ROOT = Path(__file__).parent
sys.path.insert(0, str(PROJECT_ROOT))

load_dotenv(PROJECT_ROOT / ".env")

console = Console()

def load_config() -> dict:
    """config.yaml dosyasını yükle."""
    config_path = PROJECT_ROOT / "config.yaml"
    if not config_path.exists():
        console.print("[red]HATA: config.yaml bulunamadı![/red]")
        sys.exit(1)
    with open(config_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)

def setup_logging(config: dict):
    """Loglama sistemini yapılandır."""
    log_config = config.get("logging", {})
    log_file = PROJECT_ROOT / log_config.get("file", "logs/agent_runs/latest.log")
    log_file.parent.mkdir(parents=True, exist_ok=True)
    
    logger.remove()
    logger.add(sys.stderr, level=log_config.get("level", "INFO"))
    logger.add(
        str(log_file),
        rotation=f"{log_config.get('max_size_mb', 10)} MB",
        retention=log_config.get("backup_count", 5),
        level="DEBUG",
        encoding="utf-8"
    )

@click.group()
@click.pass_context
def cli(ctx):
    """🔬 Tez Araştırma Asistanı - Otonom Araştırma Pipeline'ı"""
    ctx.ensure_object(dict)
    config = load_config()
    setup_logging(config)
    ctx.obj["config"] = config
    ctx.obj["project_root"] = PROJECT_ROOT


# ─── ADIM 9'daki tüm komutlar ───────────────────────────


@cli.command("full-scan")
@click.option("--push", is_flag=True, help="Tarama sonrasında GitHub'a push et")
@click.pass_context
def full_scan(ctx, push):
    """Tam araştırma taraması yap."""
    config = ctx.obj["config"]
    title = config.get("thesis", {}).get("title_en", config.get("project", {}).get("name", ""))
    console.print(Panel(
        f"[bold green]Tam Araştırma Pipeline Başlatılıyor[/bold green]\n"
        f"Konu: {title}",
        title="🔬 Tez Araştırma Asistanı"
    ))
    
    from agents.smart_orchestrator import SmartOrchestrator
    orchestrator = SmartOrchestrator(config, PROJECT_ROOT)
    orchestrator.run_full_pipeline()
    
    if push:
        console.print("\n[cyan]📤 GitHub'a push ediliyor...[/cyan]")
        from automation.github_sync import GitHubSync
        sync = GitHubSync(config, PROJECT_ROOT)
        sync.commit_and_push(f"Tam tarama — {datetime.now().strftime('%Y-%m-%d %H:%M')}")


@cli.command()
@click.pass_context
def incremental(ctx):
    """Günlük artımlı tarama yap."""
    config = ctx.obj["config"]
    console.print("[cyan]🔄 Artımlı tarama başlatılıyor...[/cyan]")
    
    from agents.smart_orchestrator import SmartOrchestrator
    orchestrator = SmartOrchestrator(config, PROJECT_ROOT)
    orchestrator.run_incremental()


@cli.command()
@click.pass_context
def report(ctx):
    """Haftalık rapor üret (md + docx)."""
    config = ctx.obj["config"]
    console.print("[cyan]📝 Haftalık rapor üretiliyor...[/cyan]")
    
    from agents.report_generator import ReportGenerator
    generator = ReportGenerator(config, PROJECT_ROOT)
    output = generator.generate_weekly_report()
    console.print(f"[green]✅ Rapor: {output}[/green]")


@cli.command()
@click.pass_context
def questions(ctx):
    """Danışman hocanın 3 sorusunu cevapla."""
    config = ctx.obj["config"]
    console.print(Panel(
        "[bold cyan]Danışman Sorularını Cevaplıyorum[/bold cyan]\n"
        "1. Çözülmüş ne var, çözülmemiş ne var?\n"
        "2. Kim çalışıyor, hangi gruplar, hangi şirketler?\n"
        "3. 2-3 yıl önce ile bugün arasında ne değişti?",
        title="❓ Danışman Soruları"
    ))
    
    from agents.synthesis_agent import SynthesisAgent
    synth = SynthesisAgent(config, PROJECT_ROOT)
    result = synth.run()
    
    # Sonuçları güzel göster
    if result.get("gap_analysis"):
        console.print("\n[bold green]═══ SORU 1: Çözülmüş/Çözülmemiş ═══[/bold green]")
        console.print_json(json.dumps(result["gap_analysis"], ensure_ascii=False, indent=2))
    if result.get("actor_map"):
        console.print("\n[bold green]═══ SORU 2: Kim Çalışıyor? ═══[/bold green]")
        console.print_json(json.dumps(result["actor_map"], ensure_ascii=False, indent=2))
    if result.get("trend_analysis"):
        console.print("\n[bold green]═══ SORU 3: Ne Değişti? ═══[/bold green]")
        console.print_json(json.dumps(result["trend_analysis"], ensure_ascii=False, indent=2))


@cli.command()
@click.pass_context
def proposal(ctx):
    """Tez önerisi taslağı üret."""
    config = ctx.obj["config"]
    console.print("[cyan]📋 Tez önerisi taslağı hazırlanıyor...[/cyan]")
    
    from agents.report_generator import ReportGenerator
    generator = ReportGenerator(config, PROJECT_ROOT)
    output = generator.generate_thesis_proposal()
    console.print(f"[green]✅ Tez önerisi: {output}[/green]")


@cli.command()
@click.pass_context
def daemon(ctx):
    """7/24 otomatik çalışma modunu başlat."""
    config = ctx.obj["config"]
    console.print(Panel(
        "[bold yellow]Daemon Modu Başlatılıyor[/bold yellow]\n"
        "─ Her Pazartesi 02:00 → tam tarama\n"
        "─ Her gün 08:00 → artımlı tarama\n"
        "─ Her Cuma 20:00 → haftalık rapor\n"
        "─ Her 30 dk → GitHub sync",
        title="🤖 Daemon"
    ))
    
    from automation.scheduler import ResearchScheduler
    scheduler = ResearchScheduler(config, PROJECT_ROOT)
    
    # İlk çalıştırmada hemen tam tarama yap
    console.print("[yellow]⚡ İlk tam tarama başlatılıyor...[/yellow]")
    try:
        from agents.smart_orchestrator import SmartOrchestrator
        orchestrator = SmartOrchestrator(config, PROJECT_ROOT)
        orchestrator.run_full_pipeline()
    except Exception as e:
        logger.error(f"İlk tarama hatası: {e}")
    
    scheduler.run()


@cli.command()
@click.pass_context
def status(ctx):
    """Mevcut araştırma durumunu göster."""
    config = ctx.obj["config"]
    
    table = Table(title="📊 Araştırma Durumu")
    table.add_column("Metrik", style="cyan")
    table.add_column("Değer", style="green")
    
    papers_dir = PROJECT_ROOT / "data" / "papers"
    patents_dir = PROJECT_ROOT / "data" / "patents"
    datasets_dir = PROJECT_ROOT / "data" / "datasets_catalog"
    reports_dir = PROJECT_ROOT / "reports" / "weekly"
    summaries_dir = PROJECT_ROOT / "data" / "summaries"
    
    pdf_count = len(list((papers_dir / "pdfs").glob("*.pdf"))) if (papers_dir / "pdfs").exists() else 0
    scholar_count = len(list((papers_dir / "scholar_results").glob("*.json"))) if (papers_dir / "scholar_results").exists() else 0
    patent_count = len(list(patents_dir.glob("*.json"))) if patents_dir.exists() else 0
    dataset_count = len(list(datasets_dir.glob("*.json"))) if datasets_dir.exists() else 0
    report_md = len(list(reports_dir.glob("*.md"))) if reports_dir.exists() else 0
    report_docx = len(list(reports_dir.glob("*.docx"))) if reports_dir.exists() else 0
    synthesis_count = len(list(summaries_dir.glob("*.json"))) if summaries_dir.exists() else 0
    
    table.add_row("İndirilen PDF'ler", str(pdf_count))
    table.add_row("Literatür Tarama Dosyaları", str(scholar_count))
    table.add_row("Patent Tarama Dosyaları", str(patent_count))
    table.add_row("Veri Seti Tarama Dosyaları", str(dataset_count))
    table.add_row("Sentez Dosyaları", str(synthesis_count))
    table.add_row("Raporlar (MD)", str(report_md))
    table.add_row("Raporlar (DOCX)", str(report_docx))
    
    # Son tarama bilgisi
    state_file = PROJECT_ROOT / "data" / "agent_state.json"
    if state_file.exists():
        try:
            state = json.loads(state_file.read_text())
            table.add_row("Son Tarama", state.get("last_scan", "—"))
            table.add_row("Son Rapor", state.get("last_report", "—"))
        except Exception:
            pass
    
    console.print(table)
    
    # API durumu
    api_table = Table(title="🔑 API Anahtarları")
    api_table.add_column("Servis", style="cyan")
    api_table.add_column("Durum", style="green")
    
    api_table.add_row("GEMINI_API_KEY", "✅ Ayarlanmış" if os.getenv("GEMINI_API_KEY", "").startswith("AI") else "❌ Eksik")
    api_table.add_row("SCOPUS_API_KEY", "✅ Ayarlanmış" if os.getenv("SCOPUS_API_KEY", "") else "❌ Eksik")
    api_table.add_row("GITHUB_TOKEN", "✅ Ayarlanmış" if os.getenv("GITHUB_TOKEN", "") else "❌ Eksik")
    api_table.add_row("SEMANTIC_SCHOLAR", "✅ Ayarlanmış" if os.getenv("S2_API_KEY", "") or os.getenv("SEMANTIC_SCHOLAR_API_KEY", "") else "⚠️ Opsiyonel")
    api_table.add_row("SERPAPI", "✅ Ayarlanmış" if os.getenv("SERPAPI_KEY", "") or os.getenv("SERPAPI_API_KEY", "") else "⚠️ Opsiyonel")
    api_table.add_row("TELEGRAM", "✅ Ayarlanmış" if os.getenv("TELEGRAM_BOT_TOKEN", "") else "⚠️ Opsiyonel")
    
    console.print(api_table)


@cli.command("analyze-template")
@click.pass_context
def analyze_template(ctx):
    """Rapor şablonunu analiz et."""
    config = ctx.obj["config"]
    console.print("[cyan]📄 Rapor şablonu analiz ediliyor...[/cyan]")
    
    from tools.docx_report import DocxReport
    template_path = PROJECT_ROOT / config.get("report", config.get("reporting", {})).get(
        "template_path", "templates/rapor_template.docx"
    )
    
    gen = DocxReport(template_path)
    
    try:
        from docx import Document
        doc = Document(str(template_path))
        
        table = Table(title=f"📄 Şablon Analizi: {template_path.name}")
        table.add_column("Özellik", style="cyan")
        table.add_column("Değer", style="green")
        
        table.add_row("Paragraf Sayısı", str(len(doc.paragraphs)))
        table.add_row("Tablo Sayısı", str(len(doc.tables)))
        table.add_row("Section Sayısı", str(len(doc.sections)))
        
        styles = set()
        placeholders = []
        for p in doc.paragraphs:
            styles.add(p.style.name)
            if "{{" in p.text or "<<" in p.text:
                placeholders.append(p.text)
        
        table.add_row("Kullanılan Stiller", ", ".join(sorted(styles)))
        table.add_row("Placeholder Sayısı", str(len(placeholders)))
        
        for s in doc.sections:
            has_header = any(hp.text.strip() for hp in s.header.paragraphs) if s.header else False
            has_footer = any(fp.text.strip() for fp in s.footer.paragraphs) if s.footer else False
            table.add_row("Header", "✅ Var" if has_header else "❌ Yok")
            table.add_row("Footer", "✅ Var" if has_footer else "❌ Yok")
        
        console.print(table)
        
        if placeholders:
            console.print("\n[bold]Placeholder'lar:[/bold]")
            for ph in placeholders:
                console.print(f"  → {ph}")
        
        console.print("\n[bold]İçerik Önizleme:[/bold]")
        for p in doc.paragraphs:
            if p.text.strip():
                console.print(f"  [{p.style.name}] {p.text[:100]}")
    
    except ImportError:
        console.print("[red]python-docx yüklü değil![/red]")
    except Exception as e:
        console.print(f"[red]Hata: {e}[/red]")


@cli.command("test-tools")
@click.pass_context
def test_tools(ctx):
    """Araçları test et."""
    config = ctx.obj["config"]
    console.print(Panel("[bold]🧪 Araç Testleri Başlıyor[/bold]", title="Test"))
    
    results = []
    
    # Test 1: Semantic Scholar
    try:
        from tools.semantic_scholar_api import SemanticScholarSearch
        ss = SemanticScholarSearch(config)
        r = ss.search("multi-agent EV charging", limit=3)
        results.append(("Semantic Scholar", f"✅ {len(r)} sonuç", "green"))
    except Exception as e:
        results.append(("Semantic Scholar", f"❌ {e}", "red"))
    
    # Test 2: arXiv
    try:
        from tools.arxiv_search import ArxivSearch
        ax = ArxivSearch(config)
        r = ax.search("multi-agent electric vehicle", limit=3)
        results.append(("arXiv", f"✅ {len(r)} sonuç", "green"))
    except Exception as e:
        results.append(("arXiv", f"❌ {e}", "red"))
    
    # Test 3: HuggingFace Datasets
    try:
        from tools.dataset_search import DatasetSearch
        ds = DatasetSearch(config)
        r = ds.search_huggingface("electric vehicle", max_results=3)
        results.append(("HuggingFace DS", f"✅ {len(r)} sonuç", "green"))
    except Exception as e:
        results.append(("HuggingFace DS", f"❌ {e}", "red"))
    
    # Test 4: Zenodo
    try:
        from tools.dataset_search import DatasetSearch
        ds = DatasetSearch(config)
        r = ds.search_zenodo("electric vehicle charging", max_results=3)
        results.append(("Zenodo", f"✅ {len(r)} sonuç", "green"))
    except Exception as e:
        results.append(("Zenodo", f"❌ {e}", "red"))
    
    # Test 5: PDF Parser
    try:
        from tools.pdf_parser import PDFParser
        pp = PDFParser()
        results.append(("PDF Parser", "✅ Yüklendi", "green"))
    except Exception as e:
        results.append(("PDF Parser", f"❌ {e}", "red"))
    
    # Test 6: DOCX Report
    try:
        from tools.docx_report import DocxReport
        dr = DocxReport()
        results.append(("DOCX Report", "✅ Yüklendi", "green"))
    except Exception as e:
        results.append(("DOCX Report", f"❌ {e}", "red"))
    
    # Test 7: Google Scholar
    try:
        from tools.scholar_search import ScholarSearch
        gs = ScholarSearch(config)
        results.append(("Google Scholar", "✅ Yüklendi (rate-limited)", "green"))
    except Exception as e:
        results.append(("Google Scholar", f"❌ {e}", "red"))
    
    # Test 8: Scopus
    try:
        from tools.scopus_search import ScopusSearch
        sc = ScopusSearch(config)
        has_key = "✅ API key var" if sc.api_key else "⚠️ API key yok"
        results.append(("Scopus", has_key, "green" if sc.api_key else "yellow"))
    except Exception as e:
        results.append(("Scopus", f"❌ {e}", "red"))
    
    # Test 9: LLM Interface
    try:
        from agents.llm_interface import LLMInterface
        llm = LLMInterface(config)
        has_model = "✅ Model hazır" if llm.model else "❌ Model yok"
        results.append(("LLM (Gemini)", has_model, "green" if llm.model else "red"))
    except Exception as e:
        results.append(("LLM (Gemini)", f"❌ {e}", "red"))
    
    # Sonuçları göster
    table = Table(title="🧪 Araç Test Sonuçları")
    table.add_column("Araç", style="cyan")
    table.add_column("Durum")
    for name, status_msg, color in results:
        table.add_row(name, f"[{color}]{status_msg}[/{color}]")
    console.print(table)


# ─── Eski uyumlu komutlar ────────────────────────────────


@cli.command()
@click.pass_context
def literature(ctx):
    """Sadece literatür taraması yap."""
    config = ctx.obj["config"]
    console.print("[cyan]📚 Literatür taraması başlatılıyor...[/cyan]")
    
    from agents.literature_scout import LiteratureScout
    scout = LiteratureScout(config, PROJECT_ROOT)
    scout.run()


@cli.command("patent")
@click.pass_context
def patent_cmd(ctx):
    """Sadece patent taraması yap."""
    config = ctx.obj["config"]
    console.print("[cyan]📋 Patent taraması başlatılıyor...[/cyan]")
    
    from agents.patent_scanner import PatentScanner
    scanner = PatentScanner(config, PROJECT_ROOT)
    scanner.run()


@cli.command("datasets")
@click.pass_context
def datasets_cmd(ctx):
    """Veri seti taraması yap."""
    config = ctx.obj["config"]
    console.print("[cyan]📊 Veri seti taraması başlatılıyor...[/cyan]")
    
    from agents.dataset_hunter import DatasetHunter
    hunter = DatasetHunter(config, PROJECT_ROOT)
    hunter.run()


@cli.command()
@click.pass_context
def synthesize(ctx):
    """Sentez ve gap analizi yap."""
    config = ctx.obj["config"]
    console.print("[cyan]🧪 Sentez analizi başlatılıyor...[/cyan]")
    
    from agents.synthesis_agent import SynthesisAgent
    synth = SynthesisAgent(config, PROJECT_ROOT)
    synth.run()


@cli.command("run-synthesis")
@click.argument("json_path", type=click.Path(exists=True))
@click.pass_context
def run_synthesis(ctx, json_path):
    """Sıradaki makaleyi otonom sentez sürecine al."""
    config = ctx.obj["config"]
    project_root = ctx.obj["project_root"]
    
    from agents.deep_synthesis_agent import DeepSynthesisAgent
    agent = DeepSynthesisAgent(config, project_root)
    agent.run_next(json_path)


@cli.command("run-dataset-explorer")
@click.argument("json_path", type=click.Path(exists=True))
@click.pass_context
def run_dataset_explorer(ctx, json_path):
    """Sıradaki veri setini otonom inceleme sürecine al."""
    config = ctx.obj["config"]
    project_root = ctx.obj["project_root"]
    
    from agents.dataset_explorer_agent import DatasetExplorerAgent
    agent = DatasetExplorerAgent(config, project_root)
    agent.run_next(json_path)


if __name__ == "__main__":
    cli()
